# app/services/servicio_carga_ordenes.py
# CU-03 Cargar orden de trabajo.
# Lee la hoja BOM del informe de LX, separa las órdenes por la columna A, valida, registra
# órdenes nuevas, crea versiones nuevas cuando el usuario decide reemplazar, elimina con motivo
# y pasa la orden a desarrollo (CU-04). Todo evento relevante queda en el Audit Trail.
import hashlib
import io
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload

from app.models.modelo_orden_trabajo import (
    ArchivoCargaOrden, EstadoOrden, LineaMaterialOrden, OrdenTrabajo, VersionOrdenTrabajo,
)
from app.models.modelo_usuario import Usuario
from app.services.servicio_auditoria import ServicioAuditoria

MODULO = "CU-03"
HOJA = "BOM"
# Columnas A a H de la hoja BOM, en este orden.
COLUMNAS = ["Orden", "PT", "Descripcion", "Cantidad PT", "Material", "Descripcion Material", "UM", "Cantidad"]
LETRAS = "ABCDEFGH"
MAX_ERRORES_LISTADOS = 50
MAX_BYTES_ARCHIVO = 10 * 1024 * 1024   # 10 MB: un informe de LX pesa unos pocos KB
MOTIVO_MIN, MOTIVO_MAX = 5, 500


# ====================================================================== utilidades
def _normalizar_encabezado(valor) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return " ".join(texto.lower().split())


def _texto(valor) -> str:
    """Celda a texto: 80822781 (número) y '80312190-INK' quedan ambos como texto."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return " ".join(str(valor).split())


def _decimal(valor) -> Optional[Decimal]:
    if valor is None or isinstance(valor, bool) or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        if isinstance(valor, (int, float)):
            return Decimal(str(valor))
        return Decimal(str(valor).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        return None


def _ahora() -> datetime:
    return datetime.now(timezone.utc)


def validar_motivo(motivo: Optional[str], que: str) -> str:
    texto = " ".join((motivo or "").split())
    if len(texto) < MOTIVO_MIN:
        raise MotivoInvalidoError(f"Debe escribir el motivo de {que} (mínimo {MOTIVO_MIN} caracteres).")
    if len(texto) > MOTIVO_MAX:
        raise MotivoInvalidoError(f"El motivo de {que} admite máximo {MOTIVO_MAX} caracteres.")
    return texto


# ====================================================================== resultados y errores
@dataclass
class OrdenAnalizada:
    numero_orden: str
    pt: str
    descripcion: str
    cantidad_pt: Optional[int]
    lineas: list = field(default_factory=list)      # [{"posicion","material","descripcion_material","um","cantidad"}]
    advertencias: list = field(default_factory=list)


@dataclass
class ResultadoAnalisis:
    archivo_nombre: str
    archivo_hash: str
    ordenes: list = field(default_factory=list)     # [OrdenAnalizada] en el orden en que aparecen en el archivo
    errores: list = field(default_factory=list)     # [{"regla","fila","columna","mensaje"}]
    total_errores: int = 0

    @property
    def valido(self) -> bool:
        return self.total_errores == 0

    def error(self, regla: str, mensaje: str, fila: Optional[int] = None, columna: Optional[str] = None):
        self.total_errores += 1
        if len(self.errores) < MAX_ERRORES_LISTADOS:
            self.errores.append({"regla": regla, "fila": fila, "columna": columna, "mensaje": mensaje})

    def resumen_auditoria(self) -> dict:
        return {"archivo": self.archivo_nombre, "hash": self.archivo_hash,
                "ordenes": [o.numero_orden for o in self.ordenes],
                "total_errores": self.total_errores, "errores": self.errores[:10]}


class CargaOrdenError(Exception):
    """Base de los errores de negocio de CU-03."""


class ArchivoInvalidoError(CargaOrdenError):
    def __init__(self, resultado: ResultadoAnalisis):
        self.resultado = resultado
        super().__init__("El archivo tiene errores. No se registró ninguna orden.")


class DecisionRequeridaError(CargaOrdenError):
    def __init__(self, ordenes: list):
        self.ordenes = ordenes
        super().__init__(
            "Estas órdenes ya existen y debe indicar si las reemplaza o no: " + ", ".join(ordenes) + "."
        )


class MotivoInvalidoError(CargaOrdenError):
    pass


class OrdenNoEncontradaError(CargaOrdenError):
    pass


class TransicionInvalidaError(CargaOrdenError):
    pass


# ====================================================================== servicio
class ServicioCargaOrdenes:

    # ------------------------------------------------------------------ lectura y validación
    @staticmethod
    def analizar_archivo(nombre_archivo: str, contenido: bytes) -> ResultadoAnalisis:
        """Lee la hoja BOM, separa las órdenes por la columna A y aplica las reglas de validación.
        No toca la base de datos."""
        nombre = " ".join((nombre_archivo or "").split())[:255] or "archivo_sin_nombre.xlsx"
        res = ResultadoAnalisis(archivo_nombre=nombre, archivo_hash=hashlib.sha256(contenido or b"").hexdigest())

        # VAL-01: Excel .xlsx legible y de tamaño razonable
        if not nombre.lower().endswith(".xlsx") or not contenido:
            res.error("VAL-01", "El archivo debe ser un Excel con extensión .xlsx.")
            return res
        if len(contenido) > MAX_BYTES_ARCHIVO:
            res.error("VAL-01", f"El archivo supera el tamaño máximo de {MAX_BYTES_ARCHIVO // (1024 * 1024)} MB.")
            return res
        if not contenido.startswith(b"PK"):   # todo .xlsx es un ZIP
            res.error("VAL-01", "El archivo no es un Excel válido.")
            return res
        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception as error:  # archivo dañado, protegido o que no es Excel
            res.error("VAL-01", f"No fue posible leer el archivo como Excel ({type(error).__name__}).")
            return res

        try:
            # VAL-02: hoja BOM
            hoja = next((libro[n] for n in libro.sheetnames if n.strip().upper() == HOJA), None)
            if hoja is None:
                res.error("VAL-02", "El archivo no tiene la hoja BOM.")
                return res

            filas = hoja.iter_rows(values_only=True)
            encabezado = list(next(filas, ()) or ())
            encabezado += [None] * (8 - len(encabezado))

            # VAL-03: encabezados A..H
            for i, esperado in enumerate(COLUMNAS):
                if _normalizar_encabezado(encabezado[i]) != _normalizar_encabezado(esperado):
                    res.error("VAL-03", f"El encabezado de la columna {LETRAS[i]} debe ser '{esperado}' y es "
                                        f"'{_texto(encabezado[i]) or '(vacío)'}'.", fila=1, columna=LETRAS[i])
            if not res.valido:
                return res

            # VAL-04: se ignoran las filas vacías
            datos = []
            for numero_fila, fila in enumerate(filas, start=2):
                celdas = (list(fila[:8]) + [None] * 8)[:8]
                if all(_texto(c) == "" for c in celdas):
                    continue
                datos.append((numero_fila, celdas))
        finally:
            libro.close()

        # VAL-05: al menos una fila con datos
        if not datos:
            res.error("VAL-05", "La hoja BOM no contiene órdenes.")
            return res

        # Separación de órdenes por la columna A (llave de la orden)
        ordenes: dict[str, OrdenAnalizada] = {}
        primera_fila: dict[str, int] = {}
        materiales: dict[str, dict] = {}
        for numero_fila, c in datos:
            numero = _texto(c[0])
            pt, desc, cant_pt = _texto(c[1]), _texto(c[2]), _decimal(c[3])
            material, desc_mat, um, cantidad = _texto(c[4]), _texto(c[5]), _texto(c[6]), _decimal(c[7])

            if not numero:                                                     # VAL-06
                res.error("VAL-06", "Falta el número de orden.", numero_fila, "A")
                continue
            if len(numero) > 30:
                res.error("VAL-06", "El número de orden supera 30 caracteres.", numero_fila, "A")
                continue

            if numero not in ordenes:
                ordenes[numero] = OrdenAnalizada(
                    numero_orden=numero, pt=pt, descripcion=desc,
                    cantidad_pt=int(cant_pt) if cant_pt is not None and cant_pt > 0 and cant_pt == cant_pt.to_integral_value() else None,
                )
                primera_fila[numero] = numero_fila
                materiales[numero] = {}
            orden = ordenes[numero]

            for valor, letra, campo, maximo in ((pt, "B", "PT", 50), (desc, "C", "Descripcion", 255)):
                if not valor:                                                  # VAL-10
                    res.error("VAL-10", f"La columna {letra} ({campo}) está vacía.", numero_fila, letra)
                elif len(valor) > maximo:
                    res.error("VAL-10", f"La columna {letra} ({campo}) supera {maximo} caracteres.", numero_fila, letra)

            if cant_pt is None or cant_pt <= 0 or cant_pt != cant_pt.to_integral_value():    # VAL-09
                res.error("VAL-09", "Cantidad PT debe ser un número entero mayor que cero.", numero_fila, "D")

            # VAL-08: los datos generales deben ser iguales en todas las filas de la misma orden
            ref = primera_fila[numero]
            if pt and orden.pt and pt != orden.pt:
                res.error("VAL-08", f"La orden {numero} tiene un PT distinto al de la fila {ref}.", numero_fila, "B")
            if desc and orden.descripcion and desc != orden.descripcion:
                res.error("VAL-08", f"La orden {numero} tiene una descripción distinta a la de la fila {ref}.", numero_fila, "C")
            if cant_pt is not None and orden.cantidad_pt is not None and cant_pt != orden.cantidad_pt:
                res.error("VAL-08", f"La orden {numero} tiene una Cantidad PT distinta a la de la fila {ref}.", numero_fila, "D")

            for valor, letra, campo, maximo in ((material, "E", "Material", 50), (desc_mat, "F", "Descripcion Material", 255), (um, "G", "UM", 20)):
                if not valor:                                                  # VAL-10
                    res.error("VAL-10", f"La columna {letra} ({campo}) está vacía.", numero_fila, letra)
                elif len(valor) > maximo:
                    res.error("VAL-10", f"La columna {letra} ({campo}) supera {maximo} caracteres.", numero_fila, letra)

            if cantidad is None or cantidad < 0 or cantidad >= Decimal("1e12"):   # VAL-11
                res.error("VAL-11", "Cantidad debe ser un número mayor o igual a cero.", numero_fila, "H")

            if material:                                                       # VAL-12 (advertencia)
                if material in materiales[numero]:
                    orden.advertencias.append(f"El material {material} aparece más de una vez en la orden {numero}.")
                materiales[numero][material] = numero_fila

            orden.lineas.append({"posicion": numero_fila, "material": material, "descripcion_material": desc_mat,
                                 "um": um, "cantidad": cantidad})

        res.ordenes = list(ordenes.values())
        return res

    # ------------------------------------------------------------------ consultas
    @staticmethod
    def orden_activa(db: Session, numero_orden: str, bloquear: bool = False) -> Optional[OrdenTrabajo]:
        query = db.query(OrdenTrabajo).filter(
            OrdenTrabajo.numero_orden == numero_orden, OrdenTrabajo.estado != EstadoOrden.ELIMINADA
        )
        if bloquear:
            query = query.with_for_update()
        return query.first()

    @staticmethod
    def situacion(existente: Optional[OrdenTrabajo]) -> str:
        if existente is None:
            return "NUEVA"
        return "EN_DESARROLLO" if existente.estado == EstadoOrden.EN_DESARROLLO else "EXISTE"

    @staticmethod
    def listar_ordenes(db: Session, estado: Optional[str] = None, busqueda: Optional[str] = None):
        query = db.query(OrdenTrabajo).options(joinedload(OrdenTrabajo.cargada_por), joinedload(OrdenTrabajo.eliminada_por))
        if estado == "TODAS":
            pass
        elif estado in EstadoOrden.TODOS:
            query = query.filter(OrdenTrabajo.estado == estado)
        else:  # por defecto: las órdenes vigentes
            query = query.filter(OrdenTrabajo.estado.in_(EstadoOrden.ACTIVOS))
        if busqueda:
            patron = f"%{busqueda.strip()[:60]}%"
            query = query.filter(
                OrdenTrabajo.numero_orden.ilike(patron) | OrdenTrabajo.pt.ilike(patron)
                | OrdenTrabajo.descripcion.ilike(patron)
            )
        if estado == EstadoOrden.ELIMINADA:
            return query.order_by(OrdenTrabajo.eliminada_en.desc()).limit(500).all()
        return query.order_by(OrdenTrabajo.cargada_en.desc()).limit(500).all()

    @staticmethod
    def obtener_orden(db: Session, orden_id) -> Optional[OrdenTrabajo]:
        return (
            db.query(OrdenTrabajo)
            .options(joinedload(OrdenTrabajo.lineas), joinedload(OrdenTrabajo.cargada_por), joinedload(OrdenTrabajo.creada_por))
            .filter(OrdenTrabajo.id == orden_id)
            .first()
        )

    @staticmethod
    def indicadores(db: Session) -> dict:
        conteo = dict(db.query(OrdenTrabajo.estado, func.count(OrdenTrabajo.id)).group_by(OrdenTrabajo.estado).all())
        return {
            "total": sum(conteo.values()),
            "cargadas": conteo.get(EstadoOrden.CARGADA, 0),
            "en_desarrollo": conteo.get(EstadoOrden.EN_DESARROLLO, 0),
            "eliminadas": conteo.get(EstadoOrden.ELIMINADA, 0),
        }

    # ------------------------------------------------------------------ validación previa (vista previa)
    @staticmethod
    def validar_archivo(db: Session, usuario: Usuario, nombre_archivo: str, contenido: bytes, ip_address: str):
        res = ServicioCargaOrdenes.analizar_archivo(nombre_archivo, contenido)
        if not res.valido:
            ServicioCargaOrdenes._auditar_rechazo(db, usuario, res, "VALIDAR", ip_address)
            return res, {}
        existentes = {o.numero_orden: ServicioCargaOrdenes.orden_activa(db, o.numero_orden) for o in res.ordenes}
        return res, existentes

    # ------------------------------------------------------------------ carga confirmada
    @staticmethod
    def registrar_carga(db: Session, usuario: Usuario, nombre_archivo: str, contenido: bytes,
                        decisiones: dict, ip_address: str) -> dict:
        """Registra las órdenes del archivo en una sola transacción.
        decisiones = {"<numero_orden>": {"accion": "reemplazar"|"omitir", "motivo": "..."}}
        es obligatoria solo para las órdenes que ya existen en estado CARGADA."""
        res = ServicioCargaOrdenes.analizar_archivo(nombre_archivo, contenido)
        if not res.valido:
            ServicioCargaOrdenes._auditar_rechazo(db, usuario, res, "CARGAR", ip_address)
            raise ArchivoInvalidoError(res)

        decisiones = decisiones if isinstance(decisiones, dict) else {}
        plan, pendientes = [], []
        for analizada in res.ordenes:
            existente = ServicioCargaOrdenes.orden_activa(db, analizada.numero_orden, bloquear=True)
            situacion = ServicioCargaOrdenes.situacion(existente)
            if situacion == "NUEVA":
                plan.append(("CARGAR", analizada, None, None))
            elif situacion == "EN_DESARROLLO":
                plan.append(("OMITIR", analizada, existente, "La orden ya está en desarrollo y no se puede actualizar."))
            else:
                decision = decisiones.get(analizada.numero_orden) or {}
                accion = decision.get("accion") if isinstance(decision, dict) else None
                if accion == "reemplazar":
                    motivo = validar_motivo(decision.get("motivo"), f"la actualización de la orden {analizada.numero_orden}")
                    plan.append(("ACTUALIZAR", analizada, existente, motivo))
                elif accion == "omitir":
                    plan.append(("OMITIR", analizada, existente, "El usuario decidió no reemplazarla."))
                else:
                    pendientes.append(analizada.numero_orden)
        if pendientes:
            db.rollback()
            raise DecisionRequeridaError(pendientes)

        resultado = {"cargadas": [], "actualizadas": [], "omitidas": []}
        if not any(accion != "OMITIR" for accion, *_ in plan):
            db.rollback()
            resultado["omitidas"] = [{"numero_orden": a.numero_orden, "motivo": m} for accion, a, _, m in plan]
            return resultado

        ahora = _ahora()
        eventos = []
        try:
            archivo = ArchivoCargaOrden(nombre=res.archivo_nombre, hash_sha256=res.archivo_hash, contenido=contenido,
                                        cargado_por_id=usuario.id, cargado_en=ahora)
            db.add(archivo)
            db.flush()

            for accion, analizada, existente, motivo in plan:
                if accion == "OMITIR":
                    resultado["omitidas"].append({"numero_orden": analizada.numero_orden, "motivo": motivo})
                    continue

                if accion == "CARGAR":
                    orden = OrdenTrabajo(
                        numero_orden=analizada.numero_orden, estado=EstadoOrden.CARGADA, version=1,
                        creada_por_id=usuario.id, creada_en=ahora,
                    )
                    ServicioCargaOrdenes._aplicar_datos(orden, analizada, usuario, archivo, ahora)
                    db.add(orden)
                    db.flush()
                    resultado["cargadas"].append(orden)
                    eventos.append(("ORDEN_CARGADA", "CARGAR", orden, None, ServicioCargaOrdenes._foto(orden)))
                else:  # ACTUALIZAR: nueva versión de la misma orden
                    anterior = ServicioCargaOrdenes._foto(existente)
                    db.add(VersionOrdenTrabajo(
                        orden_id=existente.id, version=existente.version, datos=anterior,
                        archivo_id=existente.archivo_id, cargada_por_id=existente.cargada_por_id,
                        cargada_en=existente.cargada_en, reemplazada_por_id=usuario.id, reemplazada_en=ahora,
                        motivo_reemplazo=motivo,
                    ))
                    existente.version += 1
                    existente.motivo_actualizacion = motivo
                    existente.lineas.clear()
                    db.flush()
                    ServicioCargaOrdenes._aplicar_datos(existente, analizada, usuario, archivo, ahora)
                    db.flush()
                    resultado["actualizadas"].append(existente)
                    eventos.append(("ORDEN_ACTUALIZADA", "ACTUALIZAR", existente, anterior,
                                    {**ServicioCargaOrdenes._foto(existente), "motivo": motivo}))
            db.commit()
        except IntegrityError:
            db.rollback()  # otra carga simultánea registró la misma orden
            raise TransicionInvalidaError("Otra persona cargó una de estas órdenes al mismo tiempo. Valide el archivo de nuevo.")
        except Exception:
            db.rollback()  # nada queda a medias
            raise

        for evento, accion, orden, antes, despues in eventos:
            ServicioAuditoria.registrar_evento(
                db=db, usuario_id=usuario.id, evento=evento, modulo=MODULO, entidad="ordenes_trabajo",
                entidad_id=orden.numero_orden, accion=accion, resultado="EXITO", ip_address=ip_address,
                datos_anteriores=antes, datos_nuevos=despues,
            )
        return resultado

    # ------------------------------------------------------------------ eliminar
    @staticmethod
    def eliminar_orden(db: Session, usuario: Usuario, orden_id, motivo: Optional[str], ip_address: str) -> OrdenTrabajo:
        orden = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == orden_id).with_for_update().first()
        if not orden:
            raise OrdenNoEncontradaError("La orden no existe.")
        if orden.estado == EstadoOrden.ELIMINADA:
            raise TransicionInvalidaError("La orden ya está eliminada.")
        if orden.estado == EstadoOrden.EN_DESARROLLO:
            raise TransicionInvalidaError("La orden está en desarrollo y ya no se puede eliminar.")
        motivo_limpio = validar_motivo(motivo, "la eliminación")

        anterior = {"estado": orden.estado, "version": orden.version}
        orden.estado = EstadoOrden.ELIMINADA
        orden.eliminada_por_id = usuario.id
        orden.eliminada_en = _ahora()
        orden.motivo_eliminacion = motivo_limpio
        db.commit()
        db.refresh(orden)
        ServicioAuditoria.registrar_evento(
            db=db, usuario_id=usuario.id, evento="ORDEN_ELIMINADA", modulo=MODULO, entidad="ordenes_trabajo",
            entidad_id=orden.numero_orden, accion="ELIMINAR", resultado="EXITO", ip_address=ip_address,
            datos_anteriores=anterior, datos_nuevos={"estado": orden.estado, "motivo": motivo_limpio},
        )
        return orden

    # ------------------------------------------------------------------ desarrollar (paso a CU-04)
    @staticmethod
    def iniciar_desarrollo(db: Session, usuario: Usuario, orden_id, ip_address: str) -> OrdenTrabajo:
        orden = db.query(OrdenTrabajo).filter(OrdenTrabajo.id == orden_id).with_for_update().first()
        if not orden:
            raise OrdenNoEncontradaError("La orden no existe.")
        if orden.estado == EstadoOrden.ELIMINADA:
            raise TransicionInvalidaError("La orden está eliminada y no se puede desarrollar.")
        if orden.estado == EstadoOrden.EN_DESARROLLO:
            db.rollback()
            return orden   # ya estaba en desarrollo: solo se continúa
        orden.estado = EstadoOrden.EN_DESARROLLO
        orden.desarrollo_por_id = usuario.id
        orden.desarrollo_en = _ahora()
        db.commit()
        db.refresh(orden)
        ServicioAuditoria.registrar_evento(
            db=db, usuario_id=usuario.id, evento="ORDEN_DESARROLLO_INICIADA", modulo=MODULO, entidad="ordenes_trabajo",
            entidad_id=orden.numero_orden, accion="DESARROLLAR", resultado="EXITO", ip_address=ip_address,
            datos_anteriores={"estado": EstadoOrden.CARGADA}, datos_nuevos={"estado": orden.estado, "version": orden.version},
        )
        return orden

    # ------------------------------------------------------------------ auditoría auxiliar
    @staticmethod
    def registrar_acceso_denegado(db: Session, usuario: Usuario, ruta: str, ip_address: str):
        ServicioAuditoria.registrar_evento(
            db=db, usuario_id=usuario.id, evento="ORDEN_ACCESO_DENEGADO", modulo=MODULO, entidad="ordenes_trabajo",
            accion="ACCESO", resultado="DENEGADO", ip_address=ip_address,
            datos_nuevos={"ruta": ruta, "rol": usuario.rol.value},
        )

    @staticmethod
    def _auditar_rechazo(db: Session, usuario: Usuario, res: ResultadoAnalisis, accion: str, ip_address: str):
        ServicioAuditoria.registrar_evento(
            db=db, usuario_id=usuario.id, evento="CARGA_ORDEN_RECHAZADA", modulo=MODULO, entidad="ordenes_trabajo",
            accion=accion, resultado="RECHAZADO", ip_address=ip_address, datos_nuevos=res.resumen_auditoria(),
        )

    # ------------------------------------------------------------------ internos
    @staticmethod
    def _aplicar_datos(orden: OrdenTrabajo, analizada: OrdenAnalizada, usuario: Usuario,
                       archivo: ArchivoCargaOrden, ahora: datetime):
        orden.pt = analizada.pt
        orden.descripcion = analizada.descripcion
        orden.cantidad_pt = analizada.cantidad_pt
        orden.total_materiales = len(analizada.lineas)
        orden.cargada_por_id = usuario.id
        orden.cargada_en = ahora
        orden.archivo_id = archivo.id
        orden.lineas = [LineaMaterialOrden(**linea) for linea in analizada.lineas]

    @staticmethod
    def _foto(orden: OrdenTrabajo) -> dict:
        return {
            "numero_orden": orden.numero_orden, "version": orden.version, "estado": orden.estado,
            "pt": orden.pt, "descripcion": orden.descripcion, "cantidad_pt": orden.cantidad_pt,
            "total_materiales": orden.total_materiales, "archivo_id": str(orden.archivo_id),
            "materiales": [
                {"material": l.material, "descripcion_material": l.descripcion_material, "um": l.um,
                 "cantidad": str(l.cantidad)} for l in orden.lineas
            ],
        }
